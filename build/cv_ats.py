# -*- coding: utf-8 -*-
"""CV completo em formato que o robo do recrutador consegue ler.

Regras seguidas aqui, todas por causa do parser do ATS:
  - uma coluna so, sem tabela, sem caixa de texto, sem cabecalho/rodape
  - nenhuma informacao dentro de imagem, icone ou grafico
  - fonte padrao (Arial no Word, Helvetica no PDF), texto preto
  - titulos de secao com os nomes que o robo procura
  - datas em MM/AAAA
  - marcador de lista escrito como caractere, nao como estilo do Word
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.colors import black
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from comum import SAIDA_CV, linha_contato, periodo

FONTE = "Arial"
BULLET = "• "


# --------------------------------------------------------------------------
# estrutura unica: as duas saidas (DOCX e PDF) leem a mesma lista de blocos
# --------------------------------------------------------------------------
def montar_blocos(d, chave_variante):
    p = d["pessoal"]
    v = d["variantes"][chave_variante]
    b = []

    b.append(("nome", p["nome"]))
    b.append(("cargo", v["cargo"]))
    b.append(("contato", linha_contato(p)))

    b.append(("secao", "RESUMO PROFISSIONAL"))
    b.append(("texto", v["resumo"]))

    b.append(("secao", "EXPERIENCIA PROFISSIONAL"))
    for exp in d["experiencia"]:
        b.append(("cargo_exp", exp["cargo"]))
        b.append(("empresa_exp", "%s | %s | %s" % (exp["empresa"], exp["local"], periodo(exp))))
        for item in exp["bullets"]:
            b.append(("bullet", item))

    b.append(("secao", "FORMACAO ACADEMICA"))
    for f in d["formacao"]:
        b.append(("cargo_exp", f["curso"]))
        b.append(("empresa_exp", "%s | %s | %s" % (f["instituicao"], f["periodo"], f["situacao"])))

    b.append(("secao", "CERTIFICACOES"))
    for c in d["certificacoes"]:
        linha = "%s | %s | %s" % (c["nome"], c["instituicao"], c["periodo"])
        if c.get("detalhe"):
            linha += " | " + c["detalhe"]
        b.append(("bullet", linha))

    b.append(("secao", "COMPETENCIAS"))
    for chave in v["competencias"]:
        g = d["competencias"][chave]
        b.append(("competencia", (g["titulo"], ", ".join(g["itens"]))))

    b.append(("secao", "IDIOMAS"))
    for i in d["idiomas"]:
        b.append(("bullet", "%s - %s" % (i["idioma"], i["nivel"])))

    return b


# --------------------------------------------------------------------------
# DOCX — e o formato que mais ATS le sem erro. Use este para upload na Gupy.
# --------------------------------------------------------------------------
def gerar_docx(d, chave_variante, caminho):
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = FONTE
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(0, 0, 0)
    pf = st.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.12

    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

    def par(texto, tamanho=10.5, negrito=False, antes=0, depois=3,
            maiusculas=False, recuo=0.0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(depois)
        if recuo:
            p.paragraph_format.left_indent = Inches(recuo)
        r = p.add_run(texto.upper() if maiusculas else texto)
        r.font.name = FONTE
        r.font.size = Pt(tamanho)
        r.bold = negrito
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    for tipo, valor in montar_blocos(d, chave_variante):
        if tipo == "nome":
            par(valor, 17, True, 0, 1)
        elif tipo == "cargo":
            par(valor, 11.5, False, 0, 1)
        elif tipo == "contato":
            par(valor, 9.5, False, 0, 8)
        elif tipo == "secao":
            par(valor, 11, True, 10, 4, maiusculas=True)
        elif tipo == "texto":
            par(valor, 10.5, False, 0, 4)
        elif tipo == "cargo_exp":
            par(valor, 11, True, 6, 0)
        elif tipo == "empresa_exp":
            par(valor, 10, False, 0, 2)
        elif tipo == "bullet":
            par(BULLET + valor, 10.5, False, 0, 2, recuo=0.16)
        elif tipo == "competencia":
            titulo, itens = valor
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            r1 = p.add_run(titulo + ": ")
            r1.bold = True
            r1.font.name = FONTE
            r1.font.size = Pt(10.5)
            r2 = p.add_run(itens)
            r2.font.name = FONTE
            r2.font.size = Pt(10.5)

    doc.save(caminho)
    return caminho


# --------------------------------------------------------------------------
# PDF — mesmo conteudo, para quando o formulario so aceita PDF
# --------------------------------------------------------------------------
def gerar_pdf(d, chave_variante, caminho):
    doc = SimpleDocTemplate(
        caminho, pagesize=A4,
        leftMargin=1.8 * cm, rightMargin=1.8 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title="Curriculo - %s" % d["pessoal"]["nome"],
        author=d["pessoal"]["nome"],
        subject=d["variantes"][chave_variante]["cargo"],
    )

    def est(nome, tam, lead, negrito=False, antes=0, depois=3, recuo=0):
        return ParagraphStyle(
            nome, fontName="Helvetica-Bold" if negrito else "Helvetica",
            fontSize=tam, leading=lead, textColor=black,
            spaceBefore=antes, spaceAfter=depois, leftIndent=recuo,
        )

    E = {
        "nome": est("nome", 17, 20, True, 0, 2),
        "cargo": est("cargo", 11.5, 14, False, 0, 2),
        "contato": est("contato", 9.5, 12, False, 0, 10),
        "secao": est("secao", 11, 14, True, 11, 5),
        "texto": est("texto", 10.5, 14.5, False, 0, 4),
        "cargo_exp": est("cargo_exp", 11, 14, True, 7, 0),
        "empresa_exp": est("empresa_exp", 10, 13, False, 0, 3),
        "bullet": est("bullet", 10.5, 14, False, 0, 3, recuo=11),
        "competencia": est("competencia", 10.5, 14, False, 0, 4),
    }

    def esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = []
    for tipo, valor in montar_blocos(d, chave_variante):
        if tipo == "competencia":
            titulo, itens = valor
            story.append(Paragraph("<b>%s:</b> %s" % (esc(titulo), esc(itens)), E["competencia"]))
        elif tipo == "bullet":
            story.append(Paragraph(BULLET + esc(valor), E["bullet"]))
        else:
            story.append(Paragraph(esc(valor), E[tipo]))

    doc.build(story)
    return caminho


def gerar_todos(d):
    base = os.path.join(SAIDA_CV, "CV_Danilo_Morelli")
    return [gerar_docx(d, "design", base + ".docx"),
            gerar_pdf(d, "design", base + ".pdf")]
